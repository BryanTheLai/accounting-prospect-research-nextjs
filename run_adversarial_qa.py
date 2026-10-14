import sqlite3
import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "leads.db"
XLSX_PATH = BASE_DIR / "master_leads.xlsx"

def run_adversarial_qa():
    print("=== PROGRAMMATIC ADVERSARIAL QA ===")
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # 1. Total count
    c.execute("SELECT COUNT(*) FROM prospects WHERE lead_id LIKE 'OC-%'")
    total_count = c.fetchone()[0]
    print(f"1. Total Outbounder Prospects Count: {total_count} (Expected: 407)")
    assert total_count == 407, f"Expected 407 records, found {total_count}"
    
    # 2. Critical null checks
    c.execute("""
    SELECT COUNT(*) FROM prospects 
    WHERE lead_id LIKE 'OC-%'
      AND (company_name IS NULL 
       OR contact_phone IS NULL 
       OR buyer_name IS NULL 
       OR root_domain IS NULL
       OR phone_script IS NULL
       OR proof_link IS NULL
       OR sg_my_wedge IS NULL)
    """)
    null_count = c.fetchone()[0]
    print(f"2. Critical NULL fields: {null_count} (Must be 0)")
    assert null_count == 0, "Critical null fields detected!"
    
    # 3. Duplicate domains
    c.execute("SELECT root_domain, COUNT(*) FROM prospects WHERE lead_id LIKE 'OC-%' GROUP BY root_domain HAVING COUNT(*) > 1")
    dup_domains = c.fetchall()
    print(f"3. Duplicate Domains: {len(dup_domains)}")
    assert len(dup_domains) == 0, f"Duplicate domains detected: {dup_domains}"
    
    # 4. Duplicate phones
    c.execute("SELECT contact_phone, COUNT(*) FROM prospects WHERE lead_id LIKE 'OC-%' GROUP BY contact_phone HAVING COUNT(*) > 1")
    dup_phones = c.fetchall()
    print(f"4. Duplicate Phones: {len(dup_phones)}")
    assert len(dup_phones) == 0, f"Duplicate phones detected: {dup_phones}"
    
    # 5. Em dash check in newly inserted records (OC-JOB-008..187 and OC-SGMY-101..220)
    c.execute("""
    SELECT lead_id, company_name FROM prospects 
    WHERE (lead_id LIKE 'OC-JOB-%' AND CAST(SUBSTR(lead_id, 8) AS INTEGER) >= 8)
       OR (lead_id LIKE 'OC-SGMY-%' AND CAST(SUBSTR(lead_id, 9) AS INTEGER) >= 101)
    """)
    new_leads = c.fetchall()
    print(f"5. Net-new verified leads verified: {len(new_leads)} (Expected: 300)")
    assert len(new_leads) == 300, f"Expected 300 net-new leads, found {len(new_leads)}"
    
    c.execute("""
    SELECT lead_id, company_name, phone_script, sg_my_wedge, notes, company_summary, bro_summary 
    FROM prospects
    WHERE (lead_id LIKE 'OC-JOB-%' AND CAST(SUBSTR(lead_id, 8) AS INTEGER) >= 8)
       OR (lead_id LIKE 'OC-SGMY-%' AND CAST(SUBSTR(lead_id, 9) AS INTEGER) >= 101)
    """)
    em_dash_violations = []
    for row in c.fetchall():
        for field in row:
            if isinstance(field, str) and "\u2014" in field:
                em_dash_violations.append((row[0], row[1]))
                break
    print(f"6. Em dash (\u2014) violations in net-new leads: {len(em_dash_violations)} (Must be 0)")
    assert len(em_dash_violations) == 0, f"Em dash violations found: {em_dash_violations}"
    
    # 7. Quality Score and Tier distribution
    c.execute("""
    SELECT 
        CASE 
            WHEN lead_id LIKE 'OC-JOB-%' AND CAST(SUBSTR(lead_id, 8) AS INTEGER) >= 8 THEN 'Archetype A (Job Hiring Pain)'
            WHEN lead_id LIKE 'OC-SGMY-%' AND CAST(SUBSTR(lead_id, 9) AS INTEGER) >= 101 THEN 'Archetype B (Cross-Border Industrial)'
            ELSE 'Historical Baseline'
        END as category,
        COUNT(*),
        AVG(quality_score),
        MIN(quality_score),
        MAX(quality_score)
    FROM prospects
    GROUP BY category
    """)
    print("\n--- Lead Quality & Archetype Distribution ---")
    for r in c.fetchall():
        print(f"Category: {r[0]} | Count: {r[1]} | Avg Score: {r[2]:.1f} | Min: {r[3]} | Max: {r[4]}")
        
    # 8. Accounting Software Distribution in Net-New Leads
    c.execute("""
    SELECT accounting_software, COUNT(*)
    FROM prospects
    WHERE (lead_id LIKE 'OC-JOB-%' AND CAST(SUBSTR(lead_id, 8) AS INTEGER) >= 8)
       OR (lead_id LIKE 'OC-SGMY-%' AND CAST(SUBSTR(lead_id, 9) AS INTEGER) >= 101)
    GROUP BY accounting_software
    ORDER BY COUNT(*) DESC
    """)
    print("\n--- Net-New Software Distribution ---")
    for r in c.fetchall():
        print(f"Software: {r[0]} | Count: {r[1]}")
        
    conn.close()
    
    # 9. Verify Excel File
    print("\n--- Verifying master_leads.xlsx ---")
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    sheet_name = 'Cross-Border Pipeline'
    assert sheet_name in wb.sheetnames, f"Sheet '{sheet_name}' not found in Excel!"
    ws = wb[sheet_name]
    excel_rows = ws.max_row
    print(f"Excel Sheet '{sheet_name}' Max Rows: {excel_rows} (1 Header + 407 Data Rows = 408)")
    assert excel_rows == 408, f"Expected 408 rows in Excel, found {excel_rows}"
    wb.close()
    
    print("\nALL ADVERSARIAL QA CHECKS PASSED PERFECTLY!")

if __name__ == "__main__":
    run_adversarial_qa()
